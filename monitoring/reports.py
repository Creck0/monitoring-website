"""
Pembuatan laporan monitoring.

Mendukung tiga format keluaran:
- CSV   : untuk diolah kembali
- XLSX  : rapi, berwarna, siap dibuka di Microsoft Excel
- PDF   : format resmi A4 siap cetak, lengkap dengan kop,
          ringkasan, analisa penyebab, dan kolom tanda tangan
"""

import csv
import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from .db import get_db, current_time

INSTANSI = "Dinas Komunikasi dan Informatika"
SUB_INSTANSI = "Sistem Monitoring Jaringan"

STATUS_COLORS = {
    "ONLINE": "16A34A",
    "SLOW": "EA580C",
    "OFFLINE": "DC2626",
    "UNKNOWN": "64748B",
}


# =========================================================
# PENGUMPULAN DATA
# =========================================================

def build_report(report_date, username="-"):
    connection = get_db()

    devices = connection.execute(
        """
        SELECT d.*, p.location AS parent_location,
               t.bottleneck_hop, t.bottleneck_ip, t.bottleneck_label
        FROM devices d
        LEFT JOIN devices p ON d.parent_id = p.id
        LEFT JOIN path_traces t ON t.device_id = d.id
        ORDER BY d.building, d.floor, d.location
        """
    ).fetchall()

    rows = []

    for device in devices:
        logs = connection.execute(
            """
            SELECT status, latency, packet_loss, jitter, checked_at
            FROM monitoring_logs
            WHERE device_id = ? AND checked_at LIKE ?
            ORDER BY checked_at ASC
            """,
            (device["id"], f"{report_date}%"),
        ).fetchall()

        total = len(logs)

        online = sum(1 for log in logs if log["status"] == "ONLINE")
        slow = sum(1 for log in logs if log["status"] == "SLOW")
        offline = sum(1 for log in logs if log["status"] == "OFFLINE")

        latencies = [
            log["latency"] for log in logs
            if log["latency"] is not None
        ]

        losses = [
            log["packet_loss"] for log in logs
            if log["packet_loss"] is not None
        ]

        place = ", ".join(
            value
            for value in (
                device["building"], device["floor"], device["room"]
            )
            if value
        )

        # Rentang waktu gangguan pada hari tersebut
        gangguan = [
            log["checked_at"][11:16]
            for log in logs
            if log["status"] == "OFFLINE"
        ]

        rows.append({
            "id": device["id"],
            "location": device["location"],
            "place": place or "-",
            "device_type": device["device_type"] or "-",
            "ip_address": device["ip_address"],
            "parent_location": device["parent_location"] or "-",
            "total": total,
            "online": online,
            "slow": slow,
            "offline": offline,
            "uptime": (
                round((online + slow) / total * 100, 2)
                if total else 0
            ),
            "avg_latency": (
                round(sum(latencies) / len(latencies), 2)
                if latencies else None
            ),
            "min_latency": round(min(latencies), 2) if latencies else None,
            "max_latency": round(max(latencies), 2) if latencies else None,
            "avg_loss": (
                round(sum(losses) / len(losses), 2) if losses else 0
            ),
            "last_status": logs[-1]["status"] if logs else "TIDAK ADA DATA",
            "last_checked": logs[-1]["checked_at"] if logs else "-",
            "first_down": gangguan[0] if gangguan else "-",
            "last_down": gangguan[-1] if gangguan else "-",
            "cause_title": device["cause_title"] or "-",
            "cause_detail": device["cause_detail"] or "-",
            "cause_scope": device["cause_scope"] or "-",
            "recommendations": json.loads(
                device["recommendations"] or "[]"
            ),
            "bottleneck_hop": device["bottleneck_hop"],
            "bottleneck_ip": device["bottleneck_ip"],
            "bottleneck_label": device["bottleneck_label"],
            "status_now": device["status"] or "UNKNOWN",
        })

    connection.close()

    total_device = len(rows)

    summary = {
        "total": total_device,
        "online": sum(1 for r in rows if r["status_now"] == "ONLINE"),
        "slow": sum(1 for r in rows if r["status_now"] == "SLOW"),
        "offline": sum(1 for r in rows if r["status_now"] == "OFFLINE"),
        "avg_uptime": (
            round(sum(r["uptime"] for r in rows) / total_device, 2)
            if total_device else 0
        ),
        "bermasalah": [
            r for r in rows
            if r["status_now"] in ("OFFLINE", "SLOW") or r["offline"] > 0
        ],
    }

    return {
        "date": report_date,
        "printed_by": username,
        "printed_at": current_time(),
        "rows": rows,
        "summary": summary,
    }


TABLE_HEADERS = [
    "No",
    "Lokasi / Perangkat",
    "Lokasi Fisik",
    "Jenis",
    "IP Address",
    "Cek",
    "Online",
    "Lambat",
    "Offline",
    "Uptime (%)",
    "Latency Rata2 (ms)",
    "Latency Min (ms)",
    "Latency Maks (ms)",
    "Packet Loss (%)",
    "Status Akhir",
    "Dugaan Penyebab",
]


def _row_values(index, row):
    return [
        index,
        row["location"],
        row["place"],
        row["device_type"],
        row["ip_address"],
        row["total"],
        row["online"],
        row["slow"],
        row["offline"],
        row["uptime"],
        row["avg_latency"] if row["avg_latency"] is not None else "-",
        row["min_latency"] if row["min_latency"] is not None else "-",
        row["max_latency"] if row["max_latency"] is not None else "-",
        row["avg_loss"],
        row["last_status"],
        row["cause_title"],
    ]


# =========================================================
# CSV
# =========================================================

def to_csv(report):
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow(["LAPORAN MONITORING JARINGAN HARIAN"])
    writer.writerow([INSTANSI])
    writer.writerow(["Tanggal", report["date"]])
    writer.writerow(["Dicetak oleh", report["printed_by"]])
    writer.writerow(["Waktu cetak", report["printed_at"]])
    writer.writerow([])

    writer.writerow(TABLE_HEADERS)

    for index, row in enumerate(report["rows"], start=1):
        writer.writerow(_row_values(index, row))

    writer.writerow([])
    writer.writerow(["ANALISA PENYEBAB DAN SARAN PENANGANAN"])
    writer.writerow([
        "Lokasi", "IP Address", "Dugaan Penyebab",
        "Penjelasan", "Saran Penanganan",
    ])

    for row in report["summary"]["bermasalah"]:
        writer.writerow([
            row["location"],
            row["ip_address"],
            row["cause_title"],
            row["cause_detail"],
            " | ".join(row["recommendations"]),
        ])

    return "\ufeff" + buffer.getvalue()


# =========================================================
# EXCEL
# =========================================================

THIN = Side(style="thin", color="D9DEE7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_title(sheet, text, subtitle, width):
    last_column = get_column_letter(width)

    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = text
    sheet["A1"].font = Font(size=15, bold=True, color="0F2A4D")
    sheet["A1"].alignment = Alignment(horizontal="center")

    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(size=11, color="45556B")
    sheet["A2"].alignment = Alignment(horizontal="center")


def to_xlsx(report):
    workbook = Workbook()

    # --- Sheet 1: Rekap ---
    sheet = workbook.active
    sheet.title = "Rekap Harian"

    width = len(TABLE_HEADERS)

    _style_title(
        sheet,
        "LAPORAN MONITORING JARINGAN HARIAN",
        f"{INSTANSI} | Tanggal {report['date']}",
        width,
    )

    sheet["A4"] = "Dicetak oleh"
    sheet["B4"] = report["printed_by"]
    sheet["A5"] = "Waktu cetak"
    sheet["B5"] = report["printed_at"]
    sheet["A6"] = "Total perangkat"
    sheet["B6"] = report["summary"]["total"]
    sheet["A7"] = "Rata-rata uptime"
    sheet["B7"] = f"{report['summary']['avg_uptime']} %"

    for row_index in range(4, 8):
        sheet.cell(row=row_index, column=1).font = Font(bold=True)

    header_row = 9

    for column_index, header in enumerate(TABLE_HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="1F4FA3")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = BORDER

    for index, row in enumerate(report["rows"], start=1):
        values = _row_values(index, row)
        excel_row = header_row + index

        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(
                row=excel_row, column=column_index, value=value
            )
            cell.border = BORDER
            cell.font = Font(size=10)
            cell.alignment = Alignment(
                horizontal="center" if column_index not in (2, 3, 16)
                else "left",
                vertical="center",
                wrap_text=column_index in (2, 3, 16),
            )

        status_cell = sheet.cell(row=excel_row, column=15)

        status_cell.font = Font(
            size=10,
            bold=True,
            color=STATUS_COLORS.get(row["last_status"], "64748B"),
        )

        if index % 2 == 0:
            for column_index in range(1, width + 1):
                sheet.cell(
                    row=excel_row, column=column_index
                ).fill = PatternFill("solid", fgColor="F6F8FB")

    widths = [5, 26, 26, 10, 15, 8, 9, 9, 9, 11, 14, 13, 14, 13, 13, 32]

    for column_index, column_width in enumerate(widths, start=1):
        sheet.column_dimensions[
            get_column_letter(column_index)
        ].width = column_width

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(width)}"
        f"{header_row + len(report['rows'])}"
    )

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    # --- Sheet 2: Analisa penyebab ---
    analysis = workbook.create_sheet("Analisa & Saran")

    _style_title(
        analysis,
        "ANALISA PENYEBAB GANGGUAN DAN SARAN PENANGANAN",
        f"Tanggal {report['date']}",
        5,
    )

    headers = [
        "Lokasi / Perangkat",
        "Lokasi Fisik",
        "IP Address",
        "Dugaan Penyebab & Penjelasan",
        "Saran Penanganan",
    ]

    for column_index, header in enumerate(headers, start=1):
        cell = analysis.cell(row=4, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="B91C1C")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = BORDER

    bermasalah = report["summary"]["bermasalah"]

    if not bermasalah:
        analysis.merge_cells("A5:E5")
        analysis["A5"] = (
            "Tidak ada perangkat yang mengalami gangguan pada "
            "tanggal ini."
        )
        analysis["A5"].alignment = Alignment(horizontal="center")

    for index, row in enumerate(bermasalah, start=5):
        values = [
            row["location"],
            row["place"],
            row["ip_address"],
            f"{row['cause_title']}\n{row['cause_detail']}",
            "\n".join(
                f"{number}. {text}"
                for number, text in enumerate(
                    row["recommendations"], start=1
                )
            ),
        ]

        for column_index, value in enumerate(values, start=1):
            cell = analysis.cell(row=index, column=column_index, value=value)
            cell.border = BORDER
            cell.font = Font(size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        analysis.row_dimensions[index].height = 90

    for column_index, column_width in enumerate([24, 24, 16, 55, 55], start=1):
        analysis.column_dimensions[
            get_column_letter(column_index)
        ].width = column_width

    analysis.page_setup.orientation = "landscape"

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return buffer


# =========================================================
# PDF
# =========================================================

def _styles():
    base = getSampleStyleSheet()

    return {
        "kop": ParagraphStyle(
            "kop", parent=base["Normal"],
            fontName="Helvetica-Bold", fontSize=14,
            alignment=TA_CENTER, textColor=colors.HexColor("#0F2A4D"),
            leading=18,
        ),
        "kop_sub": ParagraphStyle(
            "kop_sub", parent=base["Normal"],
            fontName="Helvetica", fontSize=10,
            alignment=TA_CENTER, textColor=colors.HexColor("#45556B"),
        ),
        "judul": ParagraphStyle(
            "judul", parent=base["Normal"],
            fontName="Helvetica-Bold", fontSize=12,
            alignment=TA_CENTER, spaceBefore=10, spaceAfter=2,
        ),
        "section": ParagraphStyle(
            "section", parent=base["Normal"],
            fontName="Helvetica-Bold", fontSize=11,
            textColor=colors.HexColor("#0F2A4D"),
            spaceBefore=12, spaceAfter=6,
        ),
        "body": ParagraphStyle(
            "body", parent=base["Normal"],
            fontName="Helvetica", fontSize=9,
            alignment=TA_JUSTIFY, leading=13,
        ),
        "cell": ParagraphStyle(
            "cell", parent=base["Normal"],
            fontName="Helvetica", fontSize=7.2, leading=9,
        ),
        "cell_head": ParagraphStyle(
            "cell_head", parent=base["Normal"],
            fontName="Helvetica-Bold", fontSize=7.2, leading=9,
            alignment=TA_CENTER, textColor=colors.white,
        ),
        "stat": ParagraphStyle(
            "stat", parent=base["Normal"],
            fontName="Helvetica", fontSize=7.5, leading=20,
            alignment=TA_CENTER,
        ),
        "small": ParagraphStyle(
            "small", parent=base["Normal"],
            fontName="Helvetica", fontSize=8,
            textColor=colors.HexColor("#64748B"),
        ),
    }


def _header_footer(canvas, doc, report, styles):
    canvas.saveState()

    width, height = landscape(A4)

    # Garis kop
    canvas.setStrokeColor(colors.HexColor("#1F4FA3"))
    canvas.setLineWidth(2)
    canvas.line(15 * mm, height - 28 * mm, width - 15 * mm, height - 28 * mm)

    canvas.setFillColor(colors.HexColor("#0F2A4D"))
    canvas.setFont("Helvetica-Bold", 13)
    canvas.drawCentredString(
        width / 2, height - 18 * mm, INSTANSI.upper()
    )

    canvas.setFont("Helvetica", 9)
    canvas.setFillColor(colors.HexColor("#45556B"))
    canvas.drawCentredString(
        width / 2, height - 24 * mm, SUB_INSTANSI
    )

    # Footer
    canvas.setStrokeColor(colors.HexColor("#D9DEE7"))
    canvas.setLineWidth(0.5)
    canvas.line(15 * mm, 13 * mm, width - 15 * mm, 13 * mm)

    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#64748B"))

    canvas.drawString(
        15 * mm, 9 * mm,
        f"Laporan monitoring jaringan - {report['date']} | "
        f"dicetak {report['printed_at']} oleh {report['printed_by']}",
    )

    canvas.drawRightString(
        width - 15 * mm, 9 * mm, f"Halaman {doc.page}"
    )

    canvas.restoreState()


def _status_color(status):
    return colors.HexColor("#" + STATUS_COLORS.get(status, "64748B"))


def to_pdf(report):
    styles = _styles()
    buffer = io.BytesIO()

    page_size = landscape(A4)

    doc = BaseDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=13 * mm,
        rightMargin=13 * mm,
        topMargin=32 * mm,
        bottomMargin=17 * mm,
        title=f"Laporan Monitoring Jaringan {report['date']}",
        author=INSTANSI,
    )

    frame = Frame(
        doc.leftMargin, doc.bottomMargin,
        doc.width, doc.height, id="main",
    )

    doc.addPageTemplates([
        PageTemplate(
            id="laporan",
            frames=[frame],
            onPage=lambda canvas, document: _header_footer(
                canvas, document, report, styles
            ),
        )
    ])

    story = []

    story.append(Paragraph(
        "LAPORAN HARIAN MONITORING JARINGAN", styles["judul"]
    ))

    tanggal = datetime.strptime(report["date"], "%Y-%m-%d")

    bulan = [
        "Januari", "Februari", "Maret", "April", "Mei", "Juni",
        "Juli", "Agustus", "September", "Oktober", "November",
        "Desember",
    ][tanggal.month - 1]

    story.append(Paragraph(
        f"Periode: {tanggal.day} {bulan} {tanggal.year}",
        styles["kop_sub"],
    ))

    story.append(Spacer(1, 8))

    # --- Ringkasan ---
    summary = report["summary"]

    ringkasan = Table(
        [[
            Paragraph("<b>TOTAL PERANGKAT</b><br/>"
                      f"<font size=16>{summary['total']}</font>",
                      styles["stat"]),
            Paragraph("<b>NORMAL</b><br/>"
                      f"<font size=16 color='#16A34A'>"
                      f"{summary['online']}</font>", styles["stat"]),
            Paragraph("<b>LAMBAT</b><br/>"
                      f"<font size=16 color='#EA580C'>"
                      f"{summary['slow']}</font>", styles["stat"]),
            Paragraph("<b>TIDAK TERJANGKAU</b><br/>"
                      f"<font size=16 color='#DC2626'>"
                      f"{summary['offline']}</font>", styles["stat"]),
            Paragraph("<b>RATA-RATA UPTIME</b><br/>"
                      f"<font size=16>{summary['avg_uptime']}%</font>",
                      styles["stat"]),
        ]],
        colWidths=[doc.width / 5.0] * 5,
    )

    ringkasan.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#D9DEE7")),
        ("INNERGRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#D9DEE7")),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F6F8FB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))

    story.append(ringkasan)

    # --- Tabel utama ---
    story.append(Paragraph(
        "A. Rekapitulasi Kondisi Perangkat", styles["section"]
    ))

    data = [[
        Paragraph(header, styles["cell_head"])
        for header in TABLE_HEADERS
    ]]

    for index, row in enumerate(report["rows"], start=1):
        values = _row_values(index, row)

        data.append([
            Paragraph(str(value), styles["cell"])
            for value in values
        ])

    proportions = [
        0.03, 0.115, 0.115, 0.05, 0.075, 0.035, 0.04, 0.04, 0.04,
        0.05, 0.055, 0.05, 0.055, 0.05, 0.06, 0.14,
    ]

    column_widths = [doc.width * value for value in proportions]

    table = Table(data, colWidths=column_widths, repeatRows=1)

    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4FA3")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DEE7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]

    for index, row in enumerate(report["rows"], start=1):
        if index % 2 == 0:
            style.append((
                "BACKGROUND", (0, index), (-1, index),
                colors.HexColor("#F6F8FB"),
            ))

        style.append((
            "TEXTCOLOR", (14, index), (14, index),
            _status_color(row["last_status"]),
        ))

    table.setStyle(TableStyle(style))

    story.append(table)

    if not report["rows"]:
        story.append(Paragraph(
            "Belum ada perangkat yang terdaftar.", styles["body"]
        ))

    # --- Analisa penyebab ---
    story.append(PageBreak())

    story.append(Paragraph(
        "B. Analisa Penyebab Gangguan dan Saran Penanganan",
        styles["section"],
    ))

    bermasalah = summary["bermasalah"]

    if not bermasalah:
        story.append(Paragraph(
            "Seluruh perangkat berada dalam kondisi normal pada "
            "periode ini. Tidak ditemukan gangguan yang memerlukan "
            "tindak lanjut.",
            styles["body"],
        ))

    for row in bermasalah:
        saran = "".join(
            f"<br/>{number}. {text}"
            for number, text in enumerate(
                row["recommendations"], start=1
            )
        )

        titik = (
            f"Hop {row['bottleneck_hop']} - {row['bottleneck_label']} "
            f"({row['bottleneck_ip']})"
            if row["bottleneck_ip"]
            else "Belum dilacak / tidak ditemukan lonjakan berarti"
        )

        block = Table(
            [
                [
                    Paragraph(
                        f"<b>{row['location']}</b> - {row['ip_address']}",
                        styles["body"],
                    ),
                    Paragraph(
                        f"<b>Status:</b> "
                        f"<font color='#{STATUS_COLORS.get(row['status_now'], '64748B')}'>"
                        f"{row['status_now']}</font>",
                        styles["body"],
                    ),
                ],
                [
                    Paragraph(
                        f"<b>Lokasi fisik:</b> {row['place']}<br/>"
                        f"<b>Perangkat induk:</b> "
                        f"{row['parent_location']}<br/>"
                        f"<b>Titik perlambatan:</b> {titik}<br/>"
                        f"<b>Gangguan tercatat:</b> {row['offline']} kali "
                        f"({row['first_down']} - {row['last_down']})",
                        styles["body"],
                    ),
                    Paragraph(
                        f"<b>Dugaan penyebab: {row['cause_title']}</b><br/>"
                        f"{row['cause_detail']}"
                        f"<br/><br/><b>Saran penanganan:</b>{saran}",
                        styles["body"],
                    ),
                ],
            ],
            colWidths=[doc.width * 0.38, doc.width * 0.62],
        )

        block.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#D9DEE7")),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E7EBF1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5FB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))

        story.append(KeepTogether([block, Spacer(1, 8)]))

    # --- Catatan & tanda tangan ---
    story.append(Spacer(1, 10))

    story.append(Paragraph("C. Catatan", styles["section"]))

    story.append(Paragraph(
        "Status OFFLINE berarti IP tujuan tidak membalas ICMP dari "
        "server monitoring, dan tidak selalu berarti perangkat rusak. "
        "Dugaan penyebab pada laporan ini disusun otomatis dari pola "
        "data ping (latency, packet loss, jitter, dan kondisi "
        "perangkat sekitarnya), sehingga tetap memerlukan verifikasi "
        "di lapangan sebelum tindakan perbaikan dilakukan.",
        styles["body"],
    ))

    story.append(Spacer(1, 16))

    ttd = Table(
        [
            [
                Paragraph("Mengetahui,<br/>Kepala Bidang", styles["body"]),
                "",
                Paragraph(
                    f"Malang, {tanggal.day} {bulan} {tanggal.year}<br/>"
                    "Petugas Monitoring",
                    styles["body"],
                ),
            ],
            ["", "", ""],
            [
                Paragraph(
                    "(............................................)",
                    styles["body"],
                ),
                "",
                Paragraph(
                    f"( {report['printed_by']} )", styles["body"]
                ),
            ],
        ],
        colWidths=[doc.width * 0.3, doc.width * 0.4, doc.width * 0.3],
        rowHeights=[28, 40, 18],
    )

    ttd.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))

    story.append(KeepTogether(ttd))

    doc.build(story)

    buffer.seek(0)

    return buffer
